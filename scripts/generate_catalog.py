import openpyxl
import re
import json

INPUT_XLSX = 'docs/client/catalogoV2.xlsx'
OUTPUT_CATALOG_TS = 'src/data/pim/catalog.ts'
OUTPUT_SEED_MJS = 'server/seed-data.mjs'

def clean_str(val):
    if val is None:
        return ''
    s = str(val).strip().replace('\xa0', ' ')
    replacements = {
        'FMUR': 'FÉMUR', 'TRQUEA': 'TRÁQUEA', 'CACHUPN': 'CACHUPÍN',
        'BAMB': 'BAMBÚ', 'FRMULA': 'FÓRMULA', 'CAMPAA': 'CAMPAÑA',
        'PEQUEAS': 'PEQUEÑAS', 'PEQUEA': 'PEQUEÑA', 'HOGAREO': 'HOGAREÑO',
        'CAF': 'CAFÉ', 'CARBN': 'CARBÓN', 'PULMN': 'PULMÓN',
        'DESCRIPCIN': 'DESCRIPCIÓN', 'CORAZN': 'CORAZÓN',
        'PATE': 'PATÉ', 'PLATANO': 'PLÁTANO', 'HIGADO': 'HÍGADO'
    }
    for k, v in replacements.items():
        s = s.replace(k, v)
    return s

def clean_brand(b):
    s = clean_str(b).strip()
    if s.upper() in ('N&D', 'PPVD'):
        return s.upper()
    if s.upper() in ('FIT FRMULA', 'FIT FORMULA', 'FIT FÓRMULA'):
        return 'Fit Fórmula'
    if s.upper() == '9 LIVES':
        return '9 Lives'
    if s.upper() == 'PRO PLAN':
        return 'Pro Plan'
    if s.upper() == 'DOG CHOW':
        return 'Dog Chow'
    if s.upper() == 'CAT CHOW':
        return 'Cat Chow'
    if s.upper() == 'MASTER DOG':
        return 'Master Dog'
    if s.upper() == 'MASTER CAT':
        return 'Master Cat'
    if s.upper() == 'KONGO GOLD':
        return 'Kongo Gold'
    if s.upper() == 'COOL DOG':
        return 'Cool Dog'
    if s.upper() == 'GUAU FORTE':
        return 'Guau Forte'
    if s.upper() == 'TOP ONE':
        return 'Top One'
    if s.upper() == 'FREE GO':
        return 'Free Go'
    if s.upper() == 'BIO-CAT':
        return 'Bio-Cat'
    if s.upper() == 'NATURAL MEAT':
        return 'Natural Meat'
    if s.upper() == 'STAY HAPPY':
        return 'Stay Happy'
    if s.upper() == 'TOI MOI':
        return 'Toi Moi'
    if s.upper() == 'MAXI CAT':
        return 'Maxi Cat'
    if s.upper() == 'EASY CLEAN':
        return 'Easy Clean'
    if s.upper() == 'MICHI LOVE':
        return 'Michi Love'
    return s.title() if s.isupper() else s

def clean_format(f):
    if not f: return '1 un'
    s = clean_str(f).strip().upper().replace(' ', '')
    if s.endswith('KG') or s.endswith('K'):
        num = s.replace('KG', '').replace('K', '').replace('.', ',')
        return f'{num} kg'
    if s.endswith('GR') or s.endswith('G'):
        num = s.replace('GR', '').replace('G', '')
        return f'{num} g'
    if s.endswith('ML'):
        num = s.replace('ML', '')
        return f'{num} ml'
    if s.isdigit():
        return f'{s} kg'
    if 'UNIDAD' in s or s in ('UND', 'UN'):
        return '1 un'
    if 'CAJA' in s:
        return 'Caja'
    return clean_str(f).strip()

def infer_age(text):
    t = text.upper()
    if any(k in t for k in ['CACHORRO', 'KITTEN', 'GATITO', 'PUPPY', 'JUNIOR']):
        return 'Cachorro'
    if any(k in t for k in ['SENIOR', '+7', '6+', 'ADULTO MAYOR']):
        return 'Senior'
    if any(k in t for k in ['ADULTO', 'ADULT']):
        return 'Adulto'
    return 'Todas las edades'

def calc_price(p_local, p_base):
    if p_local is not None and isinstance(p_local, (int, float)) and p_local > 0:
        return int(p_local), False
    if p_base is not None and isinstance(p_base, (int, float)) and p_base > 0:
        raw = p_base * 1.10
        rounded = int(round(raw / 1000.0) * 1000)
        return max(1000, rounded), True
    return 0, False

def build_catalog():
    wb = openpyxl.load_workbook(INPUT_XLSX)
    
    products = []
    colors = ['violet', 'green', 'peach', 'lavender']
    color_idx = 0
    alimento_img_idx = 0

    # 1. ALIMENTOS
    ws_ali = wb['Alimentos']
    curr_b = ''
    curr_p = ''
    ali_count = 0
    for r in list(ws_ali.iter_rows(values_only=True))[1:]:
        if not any(x is not None and str(x).strip() != '' for x in r): continue
        if r[0] == 'MARCA': continue
        if r[0]: curr_b = str(r[0]).strip()
        if r[1]: curr_p = str(r[1]).strip()
        
        variety = clean_str(r[2]).strip()
        formato = clean_format(r[3])
        
        p_base = r[4]
        p_local = r[5]
        p_pet = r[10] if len(r) > 10 else None
        
        effective_base = p_base if isinstance(p_base, (int, float)) else p_pet
        price, is_calc = calc_price(p_local, effective_base)
        if price <= 0:
            continue
            
        brand = clean_brand(curr_b)
        
        pet_raw = curr_p.upper()
        if 'GATO' in pet_raw:
            pet = 'Gatos'
        elif 'PERRO' in pet_raw:
            pet = 'Perros'
        else:
            pet = 'Perros Gatos'
            
        life_stage = infer_age(variety)
        
        var_title = variety.title() if variety.isupper() else variety
        if var_title.lower().startswith(brand.lower()):
            prod_name = var_title
        else:
            prod_name = f'{brand} {var_title}'
            
        detail = f'{var_title} · {formato}'
        
        badge = ''
        old_price = 0
        if str(p_local).upper() in ('PROMOCIÓN', 'PROMOCION'):
            badge = 'Promoción'
            old_price = int(round(price * 1.15 / 1000.0) * 1000)
        elif ali_count < 4:
            badge = 'SuperPrecios'
            
        img = f'/five-mascotas/alimento-{alimento_img_idx % 4}.svg'
        alimento_img_idx += 1
        
        color = colors[color_idx % len(colors)]
        color_idx += 1
        
        featured = (ali_count in [0, 1, 2, 3, 6, 8, 14, 16])
        
        products.append({
            'id': f'prod-{len(products)}',
            'name': prod_name,
            'pet': pet,
            'brand': brand,
            'category': 'Alimentos',
            'subcategory': 'Alimento seco',
            'detail': detail,
            'format': formato,
            'lifeStage': life_stage,
            'price': price,
            'oldPrice': old_price,
            'badge': badge,
            'color': color,
            'image': img,
            'featured': featured
        })
        ali_count += 1

    # 2. ARENAS
    ws_ar = wb['Arenas']
    curr_b = ''
    arena_count = 0
    for r in list(ws_ar.iter_rows(values_only=True))[1:]:
        if not any(x is not None and str(x).strip() != '' for x in r): continue
        if r[0] == 'MARCA': continue
        if r[0]: curr_b = str(r[0]).strip()
        
        aroma = clean_str(r[1]).strip()
        formato = clean_format(r[2])
        p_base = r[3]
        p_pet = r[4] if len(r) > 4 else None
        
        effective_base = p_base if isinstance(p_base, (int, float)) else None
        price, is_calc = calc_price(None, effective_base)
        if price <= 0:
            continue
            
        brand = clean_brand(curr_b)
        aroma_title = aroma.title() if aroma.isupper() else aroma
        prod_name = f'Arena {brand} {aroma_title}'
        detail = f'Arena aglomerante {aroma_title.lower()} · {formato}'
        
        badge = ''
        if p_pet and '2 X' in str(p_pet).upper():
            badge = 'Promo 2x $28.000'
        elif arena_count == 0:
            badge = 'Top Ventas'
            
        color = colors[color_idx % len(colors)]
        color_idx += 1
        
        products.append({
            'id': f'prod-{len(products)}',
            'name': prod_name,
            'pet': 'Gatos',
            'brand': brand,
            'category': 'Higiene',
            'subcategory': 'Arenas sanitarias',
            'detail': detail,
            'format': formato,
            'lifeStage': 'Todas las edades',
            'price': price,
            'oldPrice': 0,
            'badge': badge,
            'color': color,
            'image': '/demos/miga/arena.svg',
            'featured': arena_count < 2
        })
        arena_count += 1

    # 3. ANTIPARASITARIO
    ws_anti = wb['Antiparasitario']
    KNOWN_BRANDS = [
        'NEXGARD', 'BRAVECTO', 'SIMPARICA', 'CREDELIO', 'REVOLUTION',
        'ADVOCATE', 'ADVANTAGE', 'SERESTO', 'FRONTLINE', 'FIPROKILL',
        'FIPRODRAG', 'FLOVOVERMIC', 'MEVERMIC', 'DOGUIVET', 'MAMISTOP',
        'SUPERPET', 'HEPROTEC', 'HERPLEX'
    ]
    seen_anti = {}
    anti_count = 0
    for r in list(ws_anti.iter_rows(values_only=True))[1:]:
        if not any(x is not None and str(x).strip() != '' for x in r): continue
        if r[0] in ('ITEM', 'INTERNOS', None): continue
        
        item_raw = clean_str(r[0]).strip()
        p_pet = r[1]
        p_local = r[2] if len(r) > 2 else None
        
        effective_base = p_pet if isinstance(p_pet, (int, float)) else None
        price, is_calc = calc_price(p_local, effective_base)
        if price <= 0:
            continue
            
        norm_key = re.sub(r'\s+', ' ', item_raw.upper())
        if norm_key in seen_anti:
            prev_idx, prev_has_local = seen_anti[norm_key]
            if not prev_has_local and isinstance(p_local, (int, float)):
                products[prev_idx]['price'] = price
                seen_anti[norm_key] = (prev_idx, True)
            continue
            
        brand = 'FIVE Care'
        for b in KNOWN_BRANDS:
            if b in norm_key:
                brand = b.title()
                break
                
        it = norm_key
        if 'PERROS Y GATOS' in it or 'GATO Y PERRO' in it or 'PERRO Y GATO' in it:
            pet = 'Perros Gatos'
        elif any(k in it for k in ['GATO', 'GATITO', 'FELIN']):
            pet = 'Gatos'
        elif any(k in it for k in ['PERRO', 'CACHORRO', 'PUPPY', 'CANIN', 'NEXGARD', 'BRAVECTO', 'SIMPARICA', 'CREDELIO', 'DOGUIVET']):
            pet = 'Perros'
        else:
            pet = 'Perros Gatos'
            
        fmt_match = re.search(r'(\d+[\.,]?\d*\s*(?:KG|ML|MG|GR|UND|UNID|TAB|TABLETA|COMPRIMIDO|UNIDADES|UN))', item_raw, re.I)
        formato = fmt_match.group(0).lower() if fmt_match else '1 un'
        
        item_title = item_raw.title() if item_raw.isupper() else item_raw
        badge = 'Veterinario' if anti_count < 3 else ''
        
        color = colors[color_idx % len(colors)]
        color_idx += 1
        
        idx = len(products)
        products.append({
            'id': f'prod-{idx}',
            'name': item_title,
            'pet': pet,
            'brand': brand,
            'category': 'Farmacia',
            'subcategory': 'Antiparasitarios',
            'detail': f'{item_title} · Protección clínica',
            'format': formato,
            'lifeStage': 'Cachorro' if 'PUPPY' in it or 'CACHORRO' in it else 'Todas las edades',
            'price': price,
            'oldPrice': 0,
            'badge': badge,
            'color': color,
            'image': '/five-mascotas/botiquin.svg',
            'featured': anti_count < 2
        })
        seen_anti[norm_key] = (idx, isinstance(p_local, (int, float)))
        anti_count += 1

    # 4. SNACK
    ws_sn = wb['SNACK']
    curr_b = ''
    curr_p = ''
    snack_count = 0
    for r in list(ws_sn.iter_rows(values_only=True))[1:]:
        if not any(x is not None and str(x).strip() != '' for x in r): continue
        if r[0] == 'MARCA': continue
        if r[0]: curr_b = str(r[0]).strip()
        if r[1]: curr_p = str(r[1]).strip()
        
        prod_raw = clean_str(r[2]).strip()
        formato_raw = r[3]
        p_base = r[4] if len(r) > 4 else None
        
        if isinstance(formato_raw, (int, float)) and isinstance(p_base, (int, float)):
            price = int(p_base)
            formato = '1 un'
        else:
            formato = clean_format(formato_raw)
            effective_base = p_base if isinstance(p_base, (int, float)) else None
            price, is_calc = calc_price(None, effective_base)
            
        if price <= 0:
            continue
            
        brand = clean_brand(curr_b)
        pet_raw = curr_p.upper()
        if 'GATO' in pet_raw or 'GATITO' in pet_raw:
            pet = 'Gatos'
        else:
            pet = 'Perros'
            
        prod_title = prod_raw.title() if prod_raw.isupper() else prod_raw
        if prod_title.lower().startswith(brand.lower()):
            prod_name = prod_title
        else:
            prod_name = f'{brand} {prod_title}'
            
        color = colors[color_idx % len(colors)]
        color_idx += 1
        
        badge = 'Favorito' if snack_count < 2 else ''
        
        products.append({
            'id': f'prod-{len(products)}',
            'name': prod_name,
            'pet': pet,
            'brand': brand,
            'category': 'Snacks',
            'subcategory': 'Snack y premios',
            'detail': f'{prod_title} · {formato}',
            'format': formato,
            'lifeStage': 'Cachorro' if 'GATITO' in pet_raw or 'CACHORRO' in prod_raw.upper() else 'Todas las edades',
            'price': price,
            'oldPrice': 0,
            'badge': badge,
            'color': color,
            'image': '/demos/miga/snack.svg',
            'featured': snack_count < 2
        })
        snack_count += 1

    # 5. PRESERVE ACCESSORIES
    products.append({
        'id': f'prod-{len(products)}',
        'name': 'Plato de cerámica pesado antideslizante',
        'pet': 'Perros Gatos',
        'brand': 'Zee.Dog',
        'category': 'Accesorios',
        'subcategory': 'Platos y bebederos',
        'detail': '650 ml · Cerámica pesada acabado mate',
        'format': '650 ml',
        'lifeStage': 'Todas las edades',
        'price': 9990,
        'oldPrice': 11990,
        'badge': 'Oferta',
        'color': 'green',
        'image': '/demos/miga/plato.svg',
        'featured': True
    })
    products.append({
        'id': f'prod-{len(products)}',
        'name': 'Pelota con textura TPR ultra resistente',
        'pet': 'Perros',
        'brand': 'Kong',
        'category': 'Accesorios',
        'subcategory': 'Juguetes',
        'detail': 'Tamaño M · Caucho termoplástico de alta densidad',
        'format': 'Tamaño M',
        'lifeStage': 'Todas las edades',
        'price': 4990,
        'oldPrice': 0,
        'badge': '',
        'color': 'peach',
        'image': '/demos/miga/pelota.svg',
        'featured': False
    })
    products.append({
        'id': f'prod-{len(products)}',
        'name': 'Cama ortopédica viscoelástica para perro y gato',
        'pet': 'Perros Gatos',
        'brand': 'FIVE Home',
        'category': 'Accesorios',
        'subcategory': 'Camas y mantas',
        'detail': 'Espuma viscoelástica · 70x50 cm lavable',
        'format': 'Talla L',
        'lifeStage': 'Todas las edades',
        'price': 29990,
        'oldPrice': 39990,
        'badge': '2x1 Promo',
        'color': 'violet',
        'image': '/demos/miga/plato.svg',
        'featured': True
    })

    return products

def export_files(products):
    # 1. Output src/data/pim/catalog.ts
    ts_content = "import type { Product } from './types';\n\n"
    ts_content += "export const products: Product[] = [\n"
    for p in products:
        ts_content += "  {\n"
        ts_content += f"    id: {json.dumps(p['id'])},\n"
        ts_content += f"    name: {json.dumps(p['name'])},\n"
        ts_content += f"    pet: {json.dumps(p['pet'])},\n"
        ts_content += f"    brand: {json.dumps(p['brand'])},\n"
        ts_content += f"    category: {json.dumps(p['category'])},\n"
        ts_content += f"    subcategory: {json.dumps(p['subcategory'])},\n"
        ts_content += f"    detail: {json.dumps(p['detail'])},\n"
        ts_content += f"    format: {json.dumps(p['format'])},\n"
        ts_content += f"    lifeStage: {json.dumps(p['lifeStage'])},\n"
        ts_content += f"    price: {p['price']},\n"
        ts_content += f"    oldPrice: {p['oldPrice']},\n"
        if p.get('badge'):
            ts_content += f"    badge: {json.dumps(p['badge'])},\n"
        ts_content += f"    color: {json.dumps(p['color'])},\n"
        ts_content += f"    image: {json.dumps(p['image'])},\n"
        if p.get('featured'):
            ts_content += "    featured: true,\n"
        ts_content += "  },\n"
    ts_content += "];\n\n"
    ts_content += "export const formatMoney = (n: number) =>\n"
    ts_content += "  new Intl.NumberFormat('es-CL', {\n"
    ts_content += "    style: 'currency',\n"
    ts_content += "    currency: 'CLP',\n"
    ts_content += "    maximumFractionDigits: 0,\n"
    ts_content += "  }).format(n);\n"

    with open(OUTPUT_CATALOG_TS, 'w', encoding='utf-8') as f:
        f.write(ts_content)
    print(f"Wrote {len(products)} products to {OUTPUT_CATALOG_TS}")

    # 2. Update server/seed-data.mjs
    with open(OUTPUT_SEED_MJS, 'r', encoding='utf-8') as f:
        mjs_original = f.read()

    # Find where export const defaultProducts begins
    cut_idx = mjs_original.find('export const defaultProducts = [')
    if cut_idx != -1:
        mjs_prefix = mjs_original[:cut_idx]
    else:
        mjs_prefix = mjs_original

    mjs_content = mjs_prefix + "export const defaultProducts = [\n"
    for p in products:
        mjs_content += f"  {json.dumps(p, ensure_ascii=False)},\n"
    mjs_content += "];\n"

    with open(OUTPUT_SEED_MJS, 'w', encoding='utf-8') as f:
        f.write(mjs_content)
    print(f"Wrote {len(products)} products to {OUTPUT_SEED_MJS}")

if __name__ == '__main__':
    prods = build_catalog()
    export_files(prods)
