import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import csv
import re
import os

INPUT_PATH = 'docs/client/PRECIOS SEPTIEMBRE 2026.xlsx'
OUTPUT_XLSX = 'docs/client/CATALOGO_CONSOLIDADO_SEPTIEMBRE_2026.xlsx'
OUTPUT_CSV = 'docs/client/CATALOGO_CONSOLIDADO_SEPTIEMBRE_2026.csv'

def clean_str(val):
    if val is None:
        return ''
    s = str(val).strip().replace('\xa0', ' ')
    replacements = {
        'FMUR': 'FÉMUR', 'TRQUEA': 'TRÁQUEA', 'CACHUPN': 'CACHUPÍN',
        'BAMB': 'BAMBÚ', 'FRMULA': 'FÓRMULA', 'CAMPAA': 'CAMPAÑA',
        'PEQUEAS': 'PEQUEÑAS', 'PEQUEA': 'PEQUEÑA', 'HOGAREO': 'HOGAREÑO',
        'CAF': 'CAFÉ', 'CARBN': 'CARBÓN', 'PULMN': 'PULMÓN',
        'DESCRIPCIN': 'DESCRIPCIÓN', 'CORAZN': 'CORAZÓN', 'HIGADITOS': 'HIGADITOS',
        'PATE': 'PATÉ', 'PLATANO': 'PLÁTANO', 'HIGADO': 'HÍGADO'
    }
    for k, v in replacements.items():
        s = s.replace(k, v)
    return s

def clean_price(val):
    if val is None:
        return 0
    s = str(val).strip().replace('$', '').replace('.', '').replace(' ', '').replace(',', '.')
    m = re.search(r'\d+', s)
    if m:
        try:
            return int(float(m.group(0)))
        except:
            return 0
    return 0

def infer_pet(text, default='Perro'):
    t = text.upper()
    if any(k in t for k in ['GATO', 'KITTEN', 'GATITO', 'FELIN', 'CAT']):
        return 'Gato'
    if any(k in t for k in ['PERRO', 'CACHORRO', 'PUPPY', 'CANIN', 'DOG']):
        return 'Perro'
    return default

def infer_age(text):
    t = text.upper()
    if any(k in t for k in ['CACHORRO', 'KITTEN', 'GATITO', 'PUPPY', 'JUNIOR', 'BABY']):
        return 'Cachorro / Gatito'
    if any(k in t for k in ['SENIOR', '7+', '6+', 'ADULTO MAYOR', 'EDAD AVANZADA']):
        return 'Senior (+7 años)'
    if any(k in t for k in ['ADULTO', 'ADULT', 'AD ']):
        return 'Adulto'
    return 'Todas las edades'

def normalize_format(formato_raw):
    f = formato_raw.strip().upper()
    if not f:
        return 'UNIDAD'
    if re.match(r'^\d+[\.,]?\d*$', f):
        return f'{f} KG'
    f = re.sub(r'\s+', '', f)
    return f

def consolidate():
    print(f"Loading input file: {INPUT_PATH}...")
    wb_in = openpyxl.load_workbook(INPUT_PATH, data_only=True)
    supplier_sheets = [s for s in wb_in.sheetnames if not s.startswith('Hoja') and s not in ['INDICE', 'BASE']]
    
    items = []
    
    for sname in supplier_sheets:
        ws = wb_in[sname]
        prov = sname.split('.-')[-1].split(',-')[-1].strip()
        last_marca = ''
        last_mascota = ''
        
        for r in range(1, ws.max_row + 1):
            vals = [clean_str(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
            row_joined = ' '.join(vals).upper()
            
            # Skip empty or header rows
            if not row_joined:
                continue
            if 'MARCA' in row_joined and 'CATEGORIA' in row_joined:
                continue
            if 'ITEM' in row_joined and ('PRECIO PET' in row_joined or '$ PET' in row_joined):
                continue
            if 'JUGUETES' in row_joined and '$ PET' in row_joined:
                continue
            if 'ANTIPARASITARIOS PET' in row_joined:
                continue
            if 'DESCRIPCION' in row_joined and 'PRECIO NETO' in row_joined:
                continue
                
            marca = ''
            mascota = ''
            cat = 'Alimentos'
            nombre = ''
            formato = ''
            p_pet = 0
            p_local = 0
            
            if sname == '17,- ARENAS':
                cat = 'Higiene'
                mascota = 'Gato'
                marca = vals[1] if len(vals) > 1 and vals[1] else last_marca
                aroma = vals[2] if len(vals) > 2 else ''
                formato_raw = vals[3] if len(vals) > 3 else ''
                formato = normalize_format(formato_raw)
                p_local = clean_price(vals[4]) if len(vals) > 4 else 0
                if len(vals) > 5 and clean_price(vals[5]) > 0:
                    p_pet = p_local
                    p_local = clean_price(vals[5])
                else:
                    p_pet = p_local
                nombre = f'Arena Sanitaria {marca} {aroma}'.strip()
                if not nombre or p_local == 0:
                    continue
                    
            elif sname == '18,- ANTIPARASITARIOS':
                cat = 'Farmacia'
                item = vals[1] if len(vals) > 1 else ''
                if not item:
                    continue
                nombre = item
                mascota = infer_pet(item, 'Perro / Gato')
                marca = item.split()[0]
                fm = re.search(r'(\d+[\.,]?\d*\s*(?:KG|ML|MG|G|UNID|TAB|UND))', item.upper())
                formato = fm.group(0).upper() if fm else 'UNIDAD'
                p_pet = clean_price(vals[2]) if len(vals) > 2 else 0
                p_local = clean_price(vals[3]) if len(vals) > 3 else 0
                if p_local == 0 and p_pet > 0:
                    p_local = p_pet
                if p_local == 0 and p_pet == 0:
                    continue
                    
            elif sname == '19,- SNACK':
                cat = 'Snacks'
                marca = vals[1] if len(vals) > 1 and vals[1] else last_marca
                mascota = vals[2] if len(vals) > 2 and vals[2] else last_mascota
                mascota = infer_pet(mascota, 'Perro')
                nombre = vals[3] if len(vals) > 3 else ''
                formato = normalize_format(vals[4] if len(vals) > 4 and vals[4] else 'UNIDAD')
                p_local = clean_price(vals[5]) if len(vals) > 5 else 0
                p_pet = p_local
                if not nombre or p_local == 0:
                    continue
                    
            elif sname == '21,- JUGUETES':
                cat = 'Accesorios'
                mascota = 'Perro / Gato'
                nombre = vals[1] if len(vals) > 1 else ''
                marca = 'Bambú Pets' if 'BAMB' in nombre.upper() else 'Genérico'
                formato = 'UNIDAD'
                p_pet = clean_price(vals[2]) if len(vals) > 2 else 0
                p_local = clean_price(vals[3]) if len(vals) > 3 else 0
                if not nombre or p_local == 0:
                    continue
                    
            elif sname == '16.- AGROVET':
                cat = 'Alimentos'
                raw_nombre = vals[1] if len(vals) > 1 else ''
                if not raw_nombre:
                    continue
                marca = 'N&D' if 'N&D' in raw_nombre.upper() else 'Farmina'
                mascota = infer_pet(raw_nombre, 'Perro')
                fm = re.search(r'(\d+[\.,]?\d*\s*K(?:G)?)', raw_nombre, re.IGNORECASE)
                formato = fm.group(0).upper().replace('K', 'KG').replace('KGG', 'KG') if fm else '1.5KG'
                nombre = re.sub(r'\s*\d+[\.,]?\d*\s*K(?:G)?', '', raw_nombre, flags=re.IGNORECASE).strip()
                p_pet = clean_price(vals[5]) if len(vals) > 5 else 0
                p_local = clean_price(vals[6]) if len(vals) > 6 else 0
                if p_local == 0:
                    p_local = p_pet
                    
            elif sname == '10.5.- PURINA PROPLAN':
                marca = vals[1] if len(vals) > 1 and vals[1] else 'Pro Plan'
                mascota = vals[2] if len(vals) > 2 and vals[2] else last_mascota
                mascota = infer_pet(mascota, 'Perro')
                raw_nombre = vals[3] if len(vals) > 3 else ''
                fm = re.search(r'(\d+[\.,]?\d*\s*KG)', raw_nombre, re.IGNORECASE)
                formato = fm.group(0).upper() if fm else '3KG'
                nombre = re.sub(r'\s*\d+[\.,]?\d*\s*KG\s*(?:AR)?', '', raw_nombre, flags=re.IGNORECASE).strip()
                p_pet = clean_price(vals[12]) if len(vals) > 12 else 0
                p_local = clean_price(vals[13]) if len(vals) > 13 else 0
                if p_local == 0:
                    p_local = p_pet
                if not nombre or p_local == 0:
                    continue
                    
            elif sname == '22,- NUTRIQUE':
                marca = 'Nutrique'
                mascota = vals[2] if len(vals) > 2 and vals[2] else last_mascota
                raw_nombre = vals[3] if len(vals) > 3 else ''
                mascota = infer_pet(mascota or raw_nombre, 'Perro')
                clean_raw = re.sub(r'\s*\$\s*[\d\.,]+', '', raw_nombre).strip()
                fm = re.search(r'(\d+[\.,]?\d*\s*KG)', clean_raw, re.IGNORECASE)
                formato = fm.group(0).upper() if fm else 'UNIDAD'
                nombre = re.sub(r'\s*\d+[\.,]?\d*\s*KG', '', clean_raw, flags=re.IGNORECASE).strip()
                p_local = clean_price(vals[4]) if len(vals) > 4 else 0
                p_pet = p_local
                if not nombre or p_local == 0:
                    continue
                    
            else: # Standard supplier sheets
                marca = vals[1] if len(vals) > 1 and vals[1] else last_marca
                mascota = vals[2] if len(vals) > 2 and vals[2] else last_mascota
                mascota = infer_pet(mascota or vals[3], 'Perro')
                nombre = vals[3] if len(vals) > 3 else ''
                formato = normalize_format(vals[4] if len(vals) > 4 else '')
                
                p_pet = clean_price(vals[5]) if len(vals) > 5 else 0
                p_loc = clean_price(vals[6]) if len(vals) > 6 else 0
                if sname in ['1.- GEPSA', '8.- NB', '13.-ALLENDES', '11.- DRAG PHARMA']:
                    p_local = p_loc if p_loc > 0 else p_pet
                elif sname == '9.- COOPRINSEN':
                    p_local = clean_price(vals[8]) if len(vals) > 8 else p_pet
                elif sname == 'NUTRINGEN':
                    p_local = clean_price(vals[10]) if len(vals) > 10 else p_pet
                else:
                    p_local = p_pet
                    
                if not nombre or p_local == 0:
                    continue
                    
            if marca:
                last_marca = marca
            if mascota:
                last_mascota = mascota
                
            p_debito = int(round(p_local * 1.05))
            etapa = infer_age(nombre)
            
            items.append({
                'sku_id': f'FIVE-{len(items)+1:05d}',
                'proveedor_origen': prov,
                'mascota': mascota,
                'marca': (marca or last_marca or 'General').title(),
                'categoria': cat,
                'nombre_producto': nombre.strip(),
                'formato': formato.strip().upper(),
                'etapa_edad': etapa,
                'precio_pet_mayorista': p_pet,
                'precio_local_transferencia': p_local,
                'precio_local_debito': p_debito,
                'estado': 'Activo'
            })
            
    print(f"Total items parsed and normalized: {len(items)}")
    
    # 1. Export CSV with UTF-8 BOM
    headers = [
        'sku_id', 'proveedor_origen', 'mascota', 'marca', 'categoria',
        'nombre_producto', 'formato', 'etapa_edad',
        'precio_pet_mayorista', 'precio_local_transferencia', 'precio_local_debito', 'estado'
    ]
    
    os.makedirs(os.path.dirname(OUTPUT_CSV), exist_ok=True)
    with open(OUTPUT_CSV, 'w', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=headers)
        writer.writeheader()
        writer.writerows(items)
    print(f"Saved clean CSV: {OUTPUT_CSV}")
    
    # 2. Export styled XLSX workbook
    wb_out = openpyxl.Workbook()
    
    # Sheet 1: CATALOGO
    ws_cat = wb_out.active
    ws_cat.title = "CATALOGO_GENERAL"
    ws_cat.views.sheetView[0].showGridLines = True
    
    # Header styling
    header_fill = PatternFill(start_color="7025A8", end_color="7025A8", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )
    
    display_headers = [
        "SKU ID", "Proveedor Origen", "Mascota", "Marca", "Categoría",
        "Nombre Producto", "Formato", "Etapa / Edad",
        "Precio Mayorista Pet ($)", "Precio Local / Transf. ($)", "Precio Tarjeta Débito (+5%) ($)", "Estado"
    ]
    
    ws_cat.append(display_headers)
    ws_cat.row_dimensions[1].height = 28
    
    for col_idx in range(1, len(display_headers) + 1):
        cell = ws_cat.cell(1, col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border
        
    for row_idx, item in enumerate(items, start=2):
        row_vals = [
            item['sku_id'],
            item['proveedor_origen'],
            item['mascota'],
            item['marca'],
            item['categoria'],
            item['nombre_producto'],
            item['formato'],
            item['etapa_edad'],
            item['precio_pet_mayorista'],
            item['precio_local_transferencia'],
            item['precio_local_debito'],
            item['estado']
        ]
        ws_cat.append(row_vals)
        ws_cat.row_dimensions[row_idx].height = 20
        
        # Zebra striping
        zebra_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid") if row_idx % 2 == 0 else PatternFill(fill_type=None)
        
        for col_idx in range(1, len(row_vals) + 1):
            c = ws_cat.cell(row_idx, col_idx)
            c.font = data_font
            c.border = thin_border
            if zebra_fill.fill_type:
                c.fill = zebra_fill
                
            if col_idx in [1, 2, 3, 7, 8, 12]:
                c.alignment = center_align
            elif col_idx in [9, 10, 11]:
                c.alignment = right_align
                c.number_format = '$#,##0'
            else:
                c.alignment = left_align
                
    # Auto-adjust column widths
    for col in ws_cat.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_cat.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 48)
        
    # Sheet 2: RESUMEN_METRICAS
    ws_res = wb_out.create_sheet(title="RESUMEN_METRICAS")
    ws_res.views.sheetView[0].showGridLines = True
    
    # Title Banner
    ws_res.merge_cells("A1:D1")
    title_cell = ws_res.cell(1, 1, "FIVE MASCOTAS - RESUMEN DE CATÁLOGO GENERAL")
    title_cell.fill = PatternFill(start_color="1F1135", end_color="1F1135", fill_type="solid")
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    title_cell.alignment = center_align
    ws_res.row_dimensions[1].height = 35
    
    # KPI Totals
    kpis = [
        ("Total Productos Consolidados", len(items)),
        ("Categorías Únicas", len(set(i['categoria'] for i in items))),
        ("Marcas Únicas", len(set(i['marca'] for i in items))),
        ("Proveedores Relevados", len(set(i['proveedor_origen'] for i in items)))
    ]
    for idx, (label, val) in enumerate(kpis, start=3):
        ws_res.cell(idx, 1, label).font = Font(name="Calibri", size=11, bold=True)
        c_val = ws_res.cell(idx, 2, val)
        c_val.font = Font(name="Calibri", size=11, bold=True, color="7025A8")
        c_val.alignment = center_align
        
    # Breakdown by Category
    ws_res.cell(8, 1, "Categoría").font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    ws_res.cell(8, 1).fill = header_fill
    ws_res.cell(8, 2, "Cantidad SKUs").font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    ws_res.cell(8, 2).fill = header_fill
    
    cat_counts = {}
    for i in items:
        cat_counts[i['categoria']] = cat_counts.get(i['categoria'], 0) + 1
        
    r_ptr = 9
    for cat, count in sorted(cat_counts.items(), key=lambda x: x[1], reverse=True):
        ws_res.cell(r_ptr, 1, cat).font = data_font
        c_num = ws_res.cell(r_ptr, 2, count)
        c_num.font = data_font
        c_num.alignment = center_align
        r_ptr += 1
        
    # Breakdown by Species
    r_ptr += 1
    ws_res.cell(r_ptr, 1, "Especie / Mascota").font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    ws_res.cell(r_ptr, 1).fill = header_fill
    ws_res.cell(r_ptr, 2, "Cantidad SKUs").font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    ws_res.cell(r_ptr, 2).fill = header_fill
    r_ptr += 1
    
    sp_counts = {}
    for i in items:
        sp_counts[i['mascota']] = sp_counts.get(i['mascota'], 0) + 1
        
    for sp, count in sorted(sp_counts.items(), key=lambda x: x[1], reverse=True):
        ws_res.cell(r_ptr, 1, sp).font = data_font
        c_num = ws_res.cell(r_ptr, 2, count)
        c_num.font = data_font
        c_num.alignment = center_align
        r_ptr += 1
        
    for col in ws_res.columns:
        col_letter = get_column_letter(col[0].column)
        ws_res.column_dimensions[col_letter].width = 30
        
    wb_out.save(OUTPUT_XLSX)
    print(f"Saved styled Excel workbook: {OUTPUT_XLSX}")

if __name__ == '__main__':
    consolidate()
